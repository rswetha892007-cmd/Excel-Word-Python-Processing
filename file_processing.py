import pandas as pd
from docx import Document

# 1. Create sample registration data
data = {
    "Name": ["Swetha", "Arun", "Priya", "Arun", "Rahul", ""],
    "Email": [
        "SWETHA@GMAIL.COM",
        "arun@gmail.com",
        "PRIYA@GMAIL.COM",
        "arun@gmail.com",
        "rahul@gmail.com",
        " "
    ],
    "Age": [19, 20, 18, 20, 22, None],
    "Course": ["CSE", "IT", "CSE", "IT", "ECE", "CSE"]
}

df = pd.DataFrame(data)

# 2. Save the original Excel registration file
df.to_excel("registration.xlsx", index=False)

print("Original registration file created.")

# 3. Clean the Excel data
df["Name"] = df["Name"].fillna("").astype(str).str.strip()
df["Email"] = df["Email"].fillna("").astype(str).str.strip().str.lower()
df["Course"] = df["Course"].fillna("").astype(str).str.strip().str.upper()

# Remove rows with missing names or emails
df = df[(df["Name"] != "") & (df["Email"] != "")]

# Remove duplicate email registrations
df = df.drop_duplicates(subset=["Email"])

# Sort alphabetically by name
df = df.sort_values(by="Name")

# Save the cleaned Excel file
df.to_excel("cleaned_registration.xlsx", index=False)

print("Excel data cleaned and saved.")

# 4. Create a Word report
doc = Document()

doc.add_heading("Student Registration Report", level=1)
doc.add_paragraph(
    "This report contains the cleaned student registration data."
)

doc.add_paragraph(f"Total valid registrations: {len(df)}")

table = doc.add_table(rows=1, cols=4)
table.style = "Table Grid"

headers = ["Name", "Email", "Age", "Course"]

for i, header in enumerate(headers):
    table.rows[0].cells[i].text = header

for _, row in df.iterrows():
    cells = table.add_row().cells
    cells[0].text = str(row["Name"])
    cells[1].text = str(row["Email"])
    cells[2].text = (
        str(int(row["Age"]))
        if pd.notna(row["Age"])
        else "Not provided"
    )
    cells[3].text = str(row["Course"])

doc.save("registration_report.docx")

print("Word report created successfully.")
print("\nCleaned registration data:")
print(df.to_string(index=False))
print("\nAll tasks completed!")